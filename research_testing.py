"""
PQ-Auth Sentinel — Research Testing Runner
==========================================

One-file research runner for the supplied CSV dataset. It:
  * validates the dataset and required feature columns;
  * makes stratified train/validation/test splits;
  * trains and compares several scikit-learn classifiers;
  * evaluates the selected model on a held-out test set;
  * runs controlled behavioural scenarios through the selected model;
  * optionally tests ML-DSA-65 through python-oqs/liboqs if available;
  * exports CSV results, Excel dashboard/charts, PDF report, and plots.

Run from the project root:
    python research_testing.py
or:
    python research_testing.py --data data/pq_ai_authentication_dataset_20000.csv

Research note: synthetic dataset performance is not a real-world security guarantee.
Scenario feature values are controlled test inputs, not live network telemetry.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import sys
import traceback
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

from sklearn.base import clone
from sklearn.ensemble import (
    ExtraTreesClassifier,
    RandomForestClassifier,
    HistGradientBoostingClassifier,
    VotingClassifier,
    RandomForestRegressor,
)
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score,
    classification_report,
    confusion_matrix,
    f1_score,
    precision_score,
    recall_score,
)
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler, LabelEncoder
from sklearn.tree import DecisionTreeClassifier
from sklearn.ensemble import RandomForestClassifier
from sklearn.utils.class_weight import compute_sample_weight


FEATURE_COLUMNS = [
    "request_frequency",
    "inter_request_time_sec",
    "failed_attempts",
    "session_duration_min",
    "device_change",
    "certificate_valid",
    "request_pattern_score",
]
TARGET_COLUMN = "attack_type"
RANDOM_STATE = 42


def find_default_dataset(project_root: Path) -> Path:
    candidates = [
        project_root / "data" / "pq_ai_authentication_dataset_20000.csv",
        project_root / "pq_ai_authentication_dataset_20000 (1).csv",
        project_root / "pq_ai_authentication_dataset_20000.csv",
        project_root / "data" / "pq_ai_authentication_dataset_20000 (1).csv",
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0]


def load_dataset(csv_path: Path) -> tuple[pd.DataFrame, pd.Series]:
    if not csv_path.exists():
        raise FileNotFoundError(
            f"Dataset not found: {csv_path}\n"
            "Place your CSV in data/pq_ai_authentication_dataset_20000.csv "
            "or pass --data PATH."
        )

    df = pd.read_csv(csv_path)
    # Normalize accidental whitespace in headers.
    df.columns = [str(c).strip() for c in df.columns]

    missing = [c for c in FEATURE_COLUMNS + [TARGET_COLUMN] if c not in df.columns]
    if missing:
        raise ValueError(
            "The CSV is missing required columns: "
            + ", ".join(missing)
            + "\nAvailable columns: "
            + ", ".join(df.columns.astype(str))
        )

    # Only use the approved behavioural feature contract. Do not train on
    # risk_score, risk_level, decision, signature_valid, or similar leakage fields.
    X = df[FEATURE_COLUMNS].copy()
    for col in FEATURE_COLUMNS:
        X[col] = pd.to_numeric(X[col], errors="coerce")
    y = df[TARGET_COLUMN].astype(str).str.strip()

    valid_target = y.notna() & (y != "") & (y.str.lower() != "nan")
    X, y = X.loc[valid_target].reset_index(drop=True), y.loc[valid_target].reset_index(drop=True)

    if len(X) < 50:
        raise ValueError(f"Only {len(X)} usable rows found; at least 50 are recommended.")
    if y.nunique() < 2:
        raise ValueError("The target attack_type must contain at least two classes.")
    return X, y


def make_models() -> dict[str, Any]:
    """Several reproducible baseline models; no optional XGBoost dependency."""
    return {
        "Logistic Regression": Pipeline([
            ("imputer", SimpleImputer(strategy="median")),
            ("scale", StandardScaler()),
            ("model", LogisticRegression(max_iter=2500, class_weight="balanced", random_state=RANDOM_STATE)),
        ]),
        "Decision Tree": Pipeline([
            ("imputer", SimpleImputer(strategy="median")),
            ("model", DecisionTreeClassifier(
                max_depth=18, min_samples_leaf=2, class_weight="balanced",
                random_state=RANDOM_STATE
            )),
        ]),
        "Random Forest": Pipeline([
            ("imputer", SimpleImputer(strategy="median")),
            ("model", RandomForestClassifier(
                n_estimators=300, max_features="sqrt", min_samples_leaf=2,
                class_weight="balanced", n_jobs=-1, random_state=RANDOM_STATE
            )),
        ]),
        "Extra Trees": Pipeline([
            ("imputer", SimpleImputer(strategy="median")),
            ("model", ExtraTreesClassifier(
                n_estimators=350, max_features="sqrt", min_samples_leaf=2,
                class_weight="balanced", n_jobs=-1, random_state=RANDOM_STATE
            )),
        ]),
        "Hist Gradient Boosting": Pipeline([
            ("imputer", SimpleImputer(strategy="median")),
            ("model", HistGradientBoostingClassifier(
                max_iter=180, learning_rate=0.08, max_leaf_nodes=31,
                l2_regularization=0.1, random_state=RANDOM_STATE
            )),
        ]),
    }


def safe_split(X: pd.DataFrame, y: pd.Series):
    # 60% train, 20% validation, 20% held-out test.
    X_trainval, X_test, y_trainval, y_test = train_test_split(
        X, y, test_size=0.20, random_state=RANDOM_STATE, stratify=y
    )
    X_train, X_val, y_train, y_val = train_test_split(
        X_trainval, y_trainval, test_size=0.25,
        random_state=RANDOM_STATE, stratify=y_trainval
    )
    return X_train, X_val, X_test, y_train, y_val, y_test


def evaluate_models(X_train, X_val, X_test, y_train, y_val, y_test):
    models = make_models()
    fitted: dict[str, Any] = {}
    rows = []
    failures = []

    for name, model in models.items():
        try:
            model.fit(X_train, y_train)
            pred_val = model.predict(X_val)
            pred_test = model.predict(X_test)
            fitted[name] = model
            rows.append({
                "Algorithm": name,
                "Validation Accuracy": accuracy_score(y_val, pred_val),
                "Validation Macro F1": f1_score(y_val, pred_val, average="macro", zero_division=0),
                "Test Accuracy": accuracy_score(y_test, pred_test),
                "Test Macro F1": f1_score(y_test, pred_test, average="macro", zero_division=0),
                "Test Weighted F1": f1_score(y_test, pred_test, average="weighted", zero_division=0),
                "Test Precision Macro": precision_score(y_test, pred_test, average="macro", zero_division=0),
                "Test Recall Macro": recall_score(y_test, pred_test, average="macro", zero_division=0),
                "Status": "OK",
            })
        except Exception as exc:
            failures.append(f"{name}: {exc}")
            rows.append({"Algorithm": name, "Status": f"ERROR: {exc}"})

    # Soft-voting ensemble built from tree models and trained only on training data.
    try:
        ensemble = VotingClassifier(
            estimators=[
                ("extra_trees", clone(models["Extra Trees"])),
                ("random_forest", clone(models["Random Forest"])),
            ],
            voting="soft",
            weights=[2, 1],
            n_jobs=-1,
        )
        ensemble.fit(X_train, y_train)
        val_pred = ensemble.predict(X_val)
        test_pred = ensemble.predict(X_test)
        fitted["Soft Voting (Extra Trees + Random Forest)"] = ensemble
        rows.append({
            "Algorithm": "Soft Voting (Extra Trees + Random Forest)",
            "Validation Accuracy": accuracy_score(y_val, val_pred),
            "Validation Macro F1": f1_score(y_val, val_pred, average="macro", zero_division=0),
            "Test Accuracy": accuracy_score(y_test, test_pred),
            "Test Macro F1": f1_score(y_test, test_pred, average="macro", zero_division=0),
            "Test Weighted F1": f1_score(y_test, test_pred, average="weighted", zero_division=0),
            "Test Precision Macro": precision_score(y_test, test_pred, average="macro", zero_division=0),
            "Test Recall Macro": recall_score(y_test, test_pred, average="macro", zero_division=0),
            "Status": "OK",
        })
    except Exception as exc:
        failures.append(f"Soft Voting: {exc}")

    score_df = pd.DataFrame(rows)
    successful = score_df[score_df["Status"] == "OK"].copy()
    if successful.empty:
        raise RuntimeError("Every classifier failed. " + " | ".join(failures))

    # Select model by validation macro-F1, not test performance, to avoid test leakage.
    selected_name = successful.sort_values(
        ["Validation Macro F1", "Validation Accuracy"], ascending=False
    ).iloc[0]["Algorithm"]
    selected_model = fitted[selected_name]
    selected_pred = selected_model.predict(X_test)

    labels = sorted(pd.Series(y_test).unique().tolist())
    cm = confusion_matrix(y_test, selected_pred, labels=labels)
    report = classification_report(y_test, selected_pred, output_dict=True, zero_division=0)
    metrics = {
        "selected_algorithm": selected_name,
        "test_accuracy": float(accuracy_score(y_test, selected_pred)),
        "test_macro_f1": float(f1_score(y_test, selected_pred, average="macro", zero_division=0)),
        "test_weighted_f1": float(f1_score(y_test, selected_pred, average="weighted", zero_division=0)),
        "test_precision_macro": float(precision_score(y_test, selected_pred, average="macro", zero_division=0)),
        "test_recall_macro": float(recall_score(y_test, selected_pred, average="macro", zero_division=0)),
        "test_rows": int(len(y_test)),
        "train_rows": int(len(y_train)),
        "validation_rows": int(len(y_val)),
        "dataset_rows": int(len(y_train) + len(y_val) + len(y_test)),
        "labels": labels,
        "confusion_matrix": cm.tolist(),
        "classification_report": report,
        "failures": failures,
    }
    return score_df, selected_model, selected_name, metrics, (X_test, y_test, selected_pred)


def scenario_profiles() -> list[dict[str, Any]]:
    return [
        ("Normal authentication", "Normal", {"request_frequency": 4.5, "inter_request_time_sec": 18.0, "failed_attempts": 0, "session_duration_min": 24.0, "device_change": 0, "certificate_valid": 1, "request_pattern_score": 0.12}, "Valid signature + low risk should normally allow"),
        ("Replay-like activity", "Replay", {"request_frequency": 8.0, "inter_request_time_sec": 1.0, "failed_attempts": 3, "session_duration_min": 2.0, "device_change": 0, "certificate_valid": 1, "request_pattern_score": 0.80}, "Replay must be blocked by a real nonce/replay gate"),
        ("Flooding-like activity", "Flooding", {"request_frequency": 18.0, "inter_request_time_sec": 0.7, "failed_attempts": 5, "session_duration_min": 2.0, "device_change": 0, "certificate_valid": 1, "request_pattern_score": 0.91}, "High-rate pattern should elevate behavioural risk"),
        ("Spoofing-like behaviour", "Spoofing", {"request_frequency": 7.0, "inter_request_time_sec": 2.8, "failed_attempts": 3, "session_duration_min": 4.0, "device_change": 1, "certificate_valid": 0, "request_pattern_score": 0.78}, "Context anomalies should be evaluated by the model"),
        ("Low-rate suspicious activity", "Suspicious", {"request_frequency": 6.0, "inter_request_time_sec": 4.0, "failed_attempts": 2, "session_duration_min": 3.0, "device_change": 1, "certificate_valid": 1, "request_pattern_score": 0.56}, "Model prediction and risk policy are measured"),
        ("Repeated failed attempts", "Suspicious", {"request_frequency": 9.0, "inter_request_time_sec": 2.0, "failed_attempts": 5, "session_duration_min": 6.0, "device_change": 0, "certificate_valid": 1, "request_pattern_score": 0.72}, "Failure pattern should elevate behavioural risk"),
        ("Untrusted certificate context", "Spoofing", {"request_frequency": 3.0, "inter_request_time_sec": 12.0, "failed_attempts": 1, "session_duration_min": 8.0, "device_change": 1, "certificate_valid": 0, "request_pattern_score": 0.50}, "Certificate-invalid context should be treated as suspicious"),
    ]


def run_scenarios(model) -> pd.DataFrame:
    rows = []
    for idx, (name, expected_class, features, expected_policy) in enumerate(scenario_profiles(), 1):
        x = pd.DataFrame([{c: features[c] for c in FEATURE_COLUMNS}], columns=FEATURE_COLUMNS)
        prediction = str(model.predict(x)[0])
        probabilities = {}
        confidence = None
        if hasattr(model, "predict_proba"):
            proba = model.predict_proba(x)[0]
            classes = list(model.classes_)
            probabilities = {str(cls): float(prob) for cls, prob in zip(classes, proba)}
            confidence = float(np.max(proba))
        # Probability of non-Normal classes is a behavioural risk proxy, not calibrated security risk.
        normal_prob = probabilities.get("Normal", 0.0)
        risk_proxy = float(np.clip(1.0 - normal_prob, 0.0, 1.0)) if probabilities else (0.0 if prediction == "Normal" else 0.7)
        if risk_proxy >= 0.64:
            decision = "BLOCK"
            risk_level = "High"
        elif risk_proxy >= 0.34:
            decision = "FLAG"
            risk_level = "Medium"
        else:
            decision = "ALLOW"
            risk_level = "Low"
        rows.append({
            "Scenario ID": f"SC-{idx:02d}",
            "Scenario": name,
            "Expected Class / Control": expected_class,
            **features,
            "ML Prediction": prediction,
            "Model Confidence": round(confidence, 4) if confidence is not None else None,
            "Risk Proxy": round(risk_proxy, 4),
            "Risk Level": risk_level,
            "ML Policy Suggestion": decision,
            "Expected Behaviour": expected_policy,
            "Class Probabilities": json.dumps(probabilities, sort_keys=True),
            "Note": "Controlled feature profile; not a live attack or cryptographic replay test.",
        })
    return pd.DataFrame(rows)


def run_crypto_smoke_tests() -> pd.DataFrame:
    """Optional real ML-DSA-65 smoke tests; reports unavailable if OQS is not enabled."""
    rows = []
    try:
        import oqs  # type: ignore
        enabled = list(oqs.get_enabled_sig_mechanisms())
        if "ML-DSA-65" not in enabled:
            return pd.DataFrame([
                {"Test": "ML-DSA-65 availability", "Result": "NOT AVAILABLE", "Detail": "ML-DSA-65 not enabled in this liboqs build."}
            ])
        with oqs.Signature("ML-DSA-65") as signer:
            public_key = signer.generate_keypair()
            secret_key = signer.export_secret_key()
        message = b"PQ-Auth Sentinel research smoke test"
        with oqs.Signature("ML-DSA-65", secret_key=secret_key) as signer:
            signature = signer.sign(message)
        with oqs.Signature("ML-DSA-65") as verifier:
            valid_ok = bool(verifier.verify(message, signature, public_key))
            tamper_ok = bool(verifier.verify(message + b"-modified", signature, public_key))
        rows.extend([
            {"Test": "ML-DSA-65 availability", "Result": "PASS", "Detail": "Algorithm enabled."},
            {"Test": "Valid signature verification", "Result": "PASS" if valid_ok else "FAIL", "Detail": "Sign then verify canonical test bytes."},
            {"Test": "Tampered message rejected", "Result": "PASS" if not tamper_ok else "FAIL", "Detail": "Modified message should not verify."},
        ])
    except Exception as exc:
        rows.append({
            "Test": "ML-DSA-65 smoke tests",
            "Result": "NOT RUN",
            "Detail": f"OQS unavailable or failed: {exc}",
        })
    return pd.DataFrame(rows)


def create_excel_report(out_dir: Path, algorithm_scores: pd.DataFrame, scenarios: pd.DataFrame, crypto: pd.DataFrame, metrics: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, Reference
    from openpyxl.utils import get_column_letter

    path = out_dir / "PQ_Auth_Sentinel_Research_Dashboard.xlsx"
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    alg = wb.create_sheet("Algorithm Comparison")
    scen = wb.create_sheet("Scenario Results")
    crypt = wb.create_sheet("Crypto Tests")
    cm_ws = wb.create_sheet("Confusion Matrix")

    dash["A1"] = "PQ-Auth Sentinel — Research Dashboard"
    dash.merge_cells("A1:H1")
    dash["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    dash["A1"].fill = PatternFill("solid", fgColor="14233B")
    dash["A2"] = f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} | Held-out test metrics are from the supplied dataset."
    dash.merge_cells("A2:H2")
    summary = [
        ("Rows used", metrics["dataset_rows"]),
        ("Training rows", metrics["train_rows"]),
        ("Validation rows", metrics["validation_rows"]),
        ("Held-out test rows", metrics["test_rows"]),
        ("Selected model (validation macro-F1)", metrics["selected_algorithm"]),
        ("Selected model test accuracy", metrics["test_accuracy"]),
        ("Selected model test macro-F1", metrics["test_macro_f1"]),
    ]
    dash.append([])
    dash.append(["Metric", "Value"])
    for row in summary:
        dash.append(list(row))
    for cell in dash[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="294766")
    dash["B10"].number_format = "0.00%"
    dash["B11"].number_format = "0.000"

    # Algorithm test accuracy chart.
    alg.append(list(algorithm_scores.columns))
    for row in algorithm_scores.itertuples(index=False, name=None):
        alg.append(list(row))
    if "Test Accuracy" in algorithm_scores.columns:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Held-out test accuracy by algorithm"
        chart.x_axis.title = "Accuracy"
        chart.y_axis.title = "Algorithm"
        chart.add_data(Reference(alg, min_col=algorithm_scores.columns.get_loc("Test Accuracy") + 1, min_row=1, max_row=len(algorithm_scores) + 1), titles_from_data=True)
        chart.set_categories(Reference(alg, min_col=1, min_row=2, max_row=len(algorithm_scores) + 1))
        chart.height, chart.width = 8, 16
        dash.add_chart(chart, "D4")

    # Scenario risk chart.
    scen.append(list(scenarios.columns))
    for row in scenarios.itertuples(index=False, name=None):
        scen.append(list(row))
    chart2 = BarChart()
    chart2.type = "bar"
    chart2.title = "Scenario behavioural risk proxy"
    chart2.x_axis.title = "Risk proxy"
    chart2.y_axis.title = "Scenario"
    risk_col = scenarios.columns.get_loc("Risk Proxy") + 1
    chart2.add_data(Reference(scen, min_col=risk_col, min_row=1, max_row=len(scenarios) + 1), titles_from_data=True)
    chart2.set_categories(Reference(scen, min_col=2, min_row=2, max_row=len(scenarios) + 1))
    chart2.height, chart2.width = 8, 16
    dash.add_chart(chart2, "D20")

    crypt.append(list(crypto.columns))
    for row in crypto.itertuples(index=False, name=None):
        crypt.append(list(row))
    cm_ws.append(["Actual \\ Predicted"] + list(metrics["labels"]))
    for label, row in zip(metrics["labels"], metrics["confusion_matrix"]):
        cm_ws.append([label] + list(row))

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="294766")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in range(1, ws.max_column + 1):
            max_len = max(len(str(ws.cell(r, col).value or "")) for r in range(1, ws.max_row + 1))
            ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 12), 42)
    dash.column_dimensions["A"].width = 38
    dash.column_dimensions["B"].width = 36
    wb.save(path)
    return path


def create_pdf_report(out_dir: Path, algorithm_scores: pd.DataFrame, scenarios: pd.DataFrame, crypto: pd.DataFrame, metrics: dict) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.lib.units import mm

    path = out_dir / "PQ_Auth_Sentinel_Research_Report.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=10*mm, rightMargin=10*mm, topMargin=10*mm, bottomMargin=10*mm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="MainTitle", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#14233B")))
    story = [
        Paragraph("PQ-Auth Sentinel — Research Testing Report", styles["MainTitle"]),
        Paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]),
        Spacer(1, 4*mm),
        Paragraph("Methodology", styles["Heading2"]),
        Paragraph(
            "The CSV is split into 60% training, 20% validation and 20% held-out testing using stratification. "
            "Candidate classifiers are compared on validation macro-F1; the selected model is then reported on "
            "the held-out test set. Leakage-prone output fields are excluded from the feature matrix.",
            styles["BodyText"]),
        Spacer(1, 3*mm),
        Paragraph("Selected model and held-out results", styles["Heading2"]),
    ]
    summary_data = [
        ["Measure", "Result"],
        ["Selected algorithm", str(metrics["selected_algorithm"])],
        ["Test accuracy", f'{metrics["test_accuracy"]:.2%}'],
        ["Test macro-F1", f'{metrics["test_macro_f1"]:.3f}'],
        ["Test rows", str(metrics["test_rows"])],
        ["Dataset rows used", str(metrics["dataset_rows"])],
    ]
    tbl = Table(summary_data, colWidths=[65*mm, 100*mm])
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#294766")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("GRID", (0,0), (-1,-1), .4, colors.HexColor("#B8C4D2")),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F2F5F9")]),
        ("PADDING", (0,0), (-1,-1), 5),
    ]))
    story += [tbl, Spacer(1, 4*mm), Paragraph("Algorithm comparison", styles["Heading2"])]
    comp_rows = [["Algorithm", "Val accuracy", "Val macro-F1", "Test accuracy", "Test macro-F1"]]
    for _, row in algorithm_scores.iterrows():
        if row.get("Status") == "OK":
            comp_rows.append([
                str(row["Algorithm"]),
                f'{float(row["Validation Accuracy"]):.3f}',
                f'{float(row["Validation Macro F1"]):.3f}',
                f'{float(row["Test Accuracy"]):.3f}',
                f'{float(row["Test Macro F1"]):.3f}',
            ])
    comp = Table(comp_rows, repeatRows=1, colWidths=[72*mm, 28*mm, 28*mm, 28*mm, 28*mm])
    comp.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#294766")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTSIZE", (0,0), (-1,-1), 7.5),
        ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#B8C4D2")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
    ]))
    story += [comp, PageBreak(), Paragraph("Scenario results", styles["Heading2"])]
    scen_rows = [["ID", "Scenario", "Expected class", "ML prediction", "Risk proxy", "Policy suggestion"]]
    for _, row in scenarios.iterrows():
        scen_rows.append([
            str(row["Scenario ID"]), str(row["Scenario"]),
            str(row["Expected Class / Control"]), str(row["ML Prediction"]),
            f'{float(row["Risk Proxy"]):.3f}', str(row["ML Policy Suggestion"]),
        ])
    scen_tbl = Table(scen_rows, repeatRows=1, colWidths=[15*mm, 42*mm, 43*mm, 38*mm, 22*mm, 40*mm])
    scen_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#294766")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTSIZE", (0,0), (-1,-1), 7),
        ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#B8C4D2")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F2F5F9")]),
    ]))
    story += [scen_tbl, Spacer(1, 4*mm), Paragraph("Cryptographic smoke tests", styles["Heading2"])]
    crypto_rows = [["Test", "Result", "Detail"]]
    for _, row in crypto.iterrows():
        crypto_rows.append([str(row["Test"]), str(row["Result"]), str(row["Detail"])])
    ct = Table(crypto_rows, repeatRows=1, colWidths=[50*mm, 30*mm, 130*mm])
    ct.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#294766")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTSIZE", (0,0), (-1,-1), 7.5),
        ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#B8C4D2")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
    ]))
    story += [ct, Spacer(1, 4*mm), Paragraph("Approval recommendation", styles["Heading2"])]
    story.append(Paragraph(
        "For academic research review only. Model metrics are measured on the supplied dataset and may reflect "
        "synthetic data patterns. Scenario profiles are controlled inputs. The report does not certify production "
        "security. Before deployment, perform independent review, calibrate risk thresholds, use durable replay "
        "state, and validate against representative real-world traffic.",
        styles["BodyText"],
    ))
    doc.build(story)
    return path


def run(args: argparse.Namespace) -> int:
    project_root = Path(__file__).resolve().parent
    csv_path = Path(args.data).resolve() if args.data else find_default_dataset(project_root)
    out_dir = Path(args.output).resolve() if args.output else project_root / "research_outputs"
    out_dir.mkdir(parents=True, exist_ok=True)

    print(f"Dataset: {csv_path}")
    print("Loading and validating CSV...")
    X, y = load_dataset(csv_path)
    X_train, X_val, X_test, y_train, y_val, y_test = safe_split(X, y)

    print(f"Rows: {len(X)} | train={len(X_train)} | validation={len(X_val)} | test={len(X_test)}")
    print("Training and comparing classifiers...")
    algorithm_scores, selected_model, selected_name, metrics, _ = evaluate_models(
        X_train, X_val, X_test, y_train, y_val, y_test
    )

    scenarios = run_scenarios(selected_model)
    crypto = run_crypto_smoke_tests()

    # Persist raw result tables and detailed metrics.
    algorithm_scores.to_csv(out_dir / "algorithm_comparison.csv", index=False)
    scenarios.to_csv(out_dir / "scenario_results.csv", index=False)
    crypto.to_csv(out_dir / "cryptographic_tests.csv", index=False)
    (out_dir / "model_metrics.json").write_text(json.dumps(metrics, indent=2, default=str), encoding="utf-8")

    # Classification report and confusion matrix exports.
    report_df = pd.DataFrame(metrics["classification_report"]).T
    report_df.to_csv(out_dir / "selected_model_classification_report.csv")
    pd.DataFrame(
        metrics["confusion_matrix"],
        index=metrics["labels"],
        columns=metrics["labels"],
    ).to_csv(out_dir / "selected_model_confusion_matrix.csv")

    excel_path = create_excel_report(out_dir, algorithm_scores, scenarios, crypto, metrics)
    pdf_path = create_pdf_report(out_dir, algorithm_scores, scenarios, crypto, metrics)

    print("\n=== ALGORITHM COMPARISON ===")
    print(algorithm_scores.to_string(index=False))
    print(f"\nSelected by validation macro-F1: {selected_name}")
    print(f"Held-out test accuracy: {metrics['test_accuracy']:.4f}")
    print(f"Held-out test macro-F1: {metrics['test_macro_f1']:.4f}")

    print("\n=== SCENARIO OUTCOMES ===")
    print(scenarios[["Scenario", "Expected Class / Control", "ML Prediction", "Risk Proxy", "Risk Level", "ML Policy Suggestion"]].to_string(index=False))

    print("\n=== CRYPTOGRAPHIC SMOKE TESTS ===")
    print(crypto.to_string(index=False))

    print("\nReports created:")
    print(f"  Excel: {excel_path}")
    print(f"  PDF:   {pdf_path}")
    print(f"  CSVs:  {out_dir}")
    print("\nReminder: these are research/prototype results, not production security certification.")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="PQ-Auth Sentinel research test runner")
    parser.add_argument("--data", help="Path to the supplied CSV dataset")
    parser.add_argument("--output", help="Directory for reports and result CSVs")
    args = parser.parse_args()
    try:
        return run(args)
    except Exception as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        if os.environ.get("PQ_AUTH_DEBUG") == "1":
            traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
