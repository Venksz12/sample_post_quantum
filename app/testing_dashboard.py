from __future__ import annotations

from datetime import datetime, timezone
import io
from typing import Any

import pandas as pd
import streamlit as st

from app.crypto import (
    PQCryptoUnavailable,
    canonical_message,
    generate_keypair,
    sign_message,
    verify_message,
)
from app.ml_engine import ModelBundle, score_request
from app.session import (
    add_transaction,
    new_request_id,
    new_signature_id,
    new_session_id,
    now_utc,
)


def _policy(signature_valid: bool, replay_detected: bool, risk_score: float) -> tuple[str, str, float]:
    """Cryptographic and replay checks override the ML risk score."""
    if not signature_valid or replay_detected:
        return "BLOCK", "High", 1.0
    if risk_score >= 0.64:
        return "BLOCK", "High", risk_score
    if risk_score >= 0.34:
        return "FLAG", "Medium", risk_score
    return "ALLOW", "Low", risk_score


def run_scenario_suite(bundle: ModelBundle) -> pd.DataFrame:
    """
    Run controlled application-level checks using the project's real
    ML-DSA functions and loaded behavioural model. Telemetry values are
    controlled test inputs, not observations of real network traffic.
    """
    signing_key = generate_keypair()
    unrelated_key = generate_keypair()
    results: list[dict[str, Any]] = []

    cases = [
        {
            "name": "Normal authentication",
            "category": "Behavioural ML",
            "description": "Valid signature, unique nonce, normal feature profile.",
            "telemetry": {"request_frequency": 4.5, "inter_request_time_sec": 18.0, "failed_attempts": 0, "session_duration_min": 24.0, "device_change": 0, "certificate_valid": 1, "request_pattern_score": 0.12},
            "tamper": False, "corrupt_signature": False, "wrong_key": False, "replay": False,
        },
        {
            "name": "Replay attack",
            "category": "Replay protection",
            "description": "A nonce is treated as already consumed.",
            "telemetry": {"request_frequency": 8.0, "inter_request_time_sec": 1.0, "failed_attempts": 3, "session_duration_min": 2.0, "device_change": 0, "certificate_valid": 1, "request_pattern_score": 0.80},
            "tamper": False, "corrupt_signature": False, "wrong_key": False, "replay": True,
        },
        {
            "name": "Tampered message",
            "category": "ML-DSA integrity",
            "description": "The canonical message is modified after signing.",
            "telemetry": {"request_frequency": 4.0, "inter_request_time_sec": 16.0, "failed_attempts": 0, "session_duration_min": 10.0, "device_change": 0, "certificate_valid": 1, "request_pattern_score": 0.15},
            "tamper": True, "corrupt_signature": False, "wrong_key": False, "replay": False,
        },
        {
            "name": "Flooding-like behaviour",
            "category": "Behavioural ML",
            "description": "High request frequency and short intervals.",
            "telemetry": {"request_frequency": 18.0, "inter_request_time_sec": 0.7, "failed_attempts": 5, "session_duration_min": 2.0, "device_change": 0, "certificate_valid": 1, "request_pattern_score": 0.91},
            "tamper": False, "corrupt_signature": False, "wrong_key": False, "replay": False,
        },
        {
            "name": "Spoofing-like behaviour",
            "category": "Behavioural ML",
            "description": "Anomalous device/context and certificate indicators.",
            "telemetry": {"request_frequency": 7.0, "inter_request_time_sec": 2.8, "failed_attempts": 3, "session_duration_min": 4.0, "device_change": 1, "certificate_valid": 0, "request_pattern_score": 0.78},
            "tamper": False, "corrupt_signature": False, "wrong_key": False, "replay": False,
        },
        {
            "name": "Corrupted signature",
            "category": "ML-DSA integrity",
            "description": "The encoded signature is corrupted before verification.",
            "telemetry": {"request_frequency": 4.0, "inter_request_time_sec": 15.0, "failed_attempts": 0, "session_duration_min": 10.0, "device_change": 0, "certificate_valid": 1, "request_pattern_score": 0.12},
            "tamper": False, "corrupt_signature": True, "wrong_key": False, "replay": False,
        },
        {
            "name": "Wrong public key",
            "category": "ML-DSA integrity",
            "description": "Verification uses an unrelated public key.",
            "telemetry": {"request_frequency": 4.0, "inter_request_time_sec": 15.0, "failed_attempts": 0, "session_duration_min": 10.0, "device_change": 0, "certificate_valid": 1, "request_pattern_score": 0.12},
            "tamper": False, "corrupt_signature": False, "wrong_key": True, "replay": False,
        },
    ]

    for i, case in enumerate(cases, start=1):
        request_id = new_request_id()
        signature_id = new_signature_id()
        session_id = new_session_id()
        nonce = f"TEST-NONCE-{i}-{datetime.now(timezone.utc).timestamp()}"

        payload = {
            "request_id": request_id,
            "signature_id": signature_id,
            "user_id": "SCENARIO-TEST-USER",
            "application": "PQ-Auth Sentinel",
            "authentication_type": "Automated scenario test",
            "session_id": session_id,
            "timestamp": now_utc(),
            "nonce": nonce,
            "telemetry": case["telemetry"],
        }
        message = canonical_message(payload)
        signature_b64 = sign_message(message, signing_key)
        message_to_verify = message + "|TAMPERED" if case["tamper"] else message
        signature_to_verify = signature_b64

        if case["corrupt_signature"]:
            signature_to_verify = ("A" if signature_b64[:1] != "A" else "B") + signature_b64[1:]

        public_key_b64 = unrelated_key.public_key_b64 if case["wrong_key"] else signing_key.public_key_b64

        try:
            signature_valid = bool(
                verify_message(
                    message_to_verify,
                    signature_to_verify,
                    public_key_b64,
                    signing_key.algorithm,
                )
            )
        except Exception:
            signature_valid = False

        # Controlled replay test: the scenario explicitly represents a reused nonce.
        replay_detected = bool(case["replay"])

        ml_result = score_request(bundle, case["telemetry"], signature_valid)
        decision, risk_level, risk_score = _policy(
            signature_valid,
            replay_detected,
            float(ml_result.get("risk_score", 0.0)),
        )

        hard_block_case = case["replay"] or case["tamper"] or case["corrupt_signature"] or case["wrong_key"]
        if hard_block_case:
            test_status = "PASS" if decision == "BLOCK" else "FAIL"
            expected = "BLOCK"
        elif case["name"] == "Normal authentication":
            # Valid cryptography is the expected result; the behavioural model may flag it.
            test_status = "PASS" if signature_valid and not replay_detected else "FAIL"
            expected = "Valid signature; ML decides access"
        else:
            # These tests exercise classifier response; do not hard-code a guaranteed label.
            test_status = "PASS" if signature_valid else "FAIL"
            expected = "Valid signature; inspect ML risk decision"

        results.append({
            "Scenario ID": f"SC-{i:02d}",
            "Scenario": case["name"],
            "Category": case["category"],
            "Description": case["description"],
            "Signature Valid": signature_valid,
            "Replay Detected": replay_detected,
            "ML Prediction": ml_result.get("attack_type", "Unknown"),
            "Risk Score": round(float(risk_score), 4),
            "Risk Level": risk_level,
            "Decision": decision,
            "Expected Policy": expected,
            "Test Status": test_status,
            "Request ID": request_id,
            "Signature ID": signature_id,
        })

        # Session ledger record; private/secret key is never included.
        try:
            add_transaction({
                **payload,
                "algorithm": signing_key.algorithm,
                "key_fingerprint": signing_key.key_fingerprint,
                "message": message_to_verify,
                "signature": signature_to_verify,
                "public_key": public_key_b64,
                "signature_valid": signature_valid,
                "signature_status": "VALID" if signature_valid else "INVALID",
                "replay_detected": replay_detected,
                "attack_type": ml_result.get("attack_type", "Unknown"),
                "risk_score": risk_score,
                "risk_level": risk_level,
                "decision": decision,
                "probability_risk": ml_result.get("probability_risk", risk_score),
                "class_probabilities": ml_result.get("class_probabilities", {}),
                "tampered": bool(case["tamper"] or case["corrupt_signature"] or case["wrong_key"]),
                "test_scenario": case["name"],
            })
        except Exception:
            # Reporting should still complete if the session ledger has a stricter schema.
            pass

    return pd.DataFrame(results)


def build_excel_report(results: pd.DataFrame, model_metrics: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, Reference
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"
    detail = wb.create_sheet("Scenario Results")
    coverage = wb.create_sheet("Algorithm Coverage")

    dashboard["A1"] = "PQ-Auth Sentinel — Scenario Test Dashboard"
    dashboard.merge_cells("A1:H1")
    dashboard["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    dashboard["A1"].fill = PatternFill("solid", fgColor="14233B")
    dashboard["A2"] = "Generated from the current controlled test run. Prototype evidence only."
    dashboard.merge_cells("A2:H2")

    total = len(results)
    passed = int((results["Test Status"] == "PASS").sum())
    failed = total - passed
    metrics = [
        ("Scenarios executed", total),
        ("Checks passed", passed),
        ("Checks failed", failed),
        ("Pass rate", passed / total if total else 0),
        ("ALLOW decisions", int((results["Decision"] == "ALLOW").sum())),
        ("FLAG decisions", int((results["Decision"] == "FLAG").sum())),
        ("BLOCK decisions", int((results["Decision"] == "BLOCK").sum())),
        ("Model test accuracy (synthetic data)", model_metrics.get("test_accuracy", "N/A")),
        ("Model macro-F1 (synthetic data)", model_metrics.get("test_macro_f1", "N/A")),
    ]
    dashboard.append([])
    dashboard.append(["Metric", "Value"])
    for row in metrics:
        dashboard.append(list(row))
    for cell in dashboard[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="294766")
    dashboard["B8"].number_format = "0.0%"

    # Risk by scenario chart data
    dashboard["D4"] = "Scenario"
    dashboard["E4"] = "Risk Score"
    for row_idx, (_, row) in enumerate(results.iterrows(), start=5):
        dashboard.cell(row_idx, 4, row["Scenario"])
        dashboard.cell(row_idx, 5, float(row["Risk Score"]))
    risk_chart = BarChart()
    risk_chart.type = "bar"
    risk_chart.title = "Risk score by scenario"
    risk_chart.y_axis.title = "Scenario"
    risk_chart.x_axis.title = "Risk score"
    risk_chart.add_data(Reference(dashboard, min_col=5, min_row=4, max_row=4 + total), titles_from_data=True)
    risk_chart.set_categories(Reference(dashboard, min_col=4, min_row=5, max_row=4 + total))
    risk_chart.height = 7
    risk_chart.width = 15
    dashboard.add_chart(risk_chart, "G4")

    dashboard["A15"] = "Decision"
    dashboard["B15"] = "Count"
    for r, decision in enumerate(("ALLOW", "FLAG", "BLOCK"), start=16):
        dashboard.cell(r, 1, decision)
        dashboard.cell(r, 2, int((results["Decision"] == decision).sum()))
    decision_chart = BarChart()
    decision_chart.type = "col"
    decision_chart.title = "Decision distribution"
    decision_chart.y_axis.title = "Count"
    decision_chart.x_axis.title = "Decision"
    decision_chart.add_data(Reference(dashboard, min_col=2, min_row=15, max_row=18), titles_from_data=True)
    decision_chart.set_categories(Reference(dashboard, min_col=1, min_row=16, max_row=18))
    decision_chart.height = 6
    decision_chart.width = 10
    dashboard.add_chart(decision_chart, "D20")

    detail.append(list(results.columns))
    for row in results.itertuples(index=False, name=None):
        detail.append(list(row))

    coverage.append(["Area", "Algorithm/control", "What is tested", "Limitation"])
    for row in [
        ["Digital signatures", "ML-DSA-65", "Valid message, tampering, corrupted signature, wrong public key", "Requires working liboqs ML-DSA-65 implementation"],
        ["Replay protection", "Nonce/replay policy", "Controlled reused-nonce case must be blocked", "This scenario suite simulates replay state; persistence is session-only unless separately implemented"],
        ["Behavioural classification", "Existing voting classifier", "Classifies seven telemetry features", "Training data is synthetic; results are not real-world accuracy guarantees"],
        ["Risk policy", "Hard gates + score thresholds", "Signature/replay block; threshold-based allow/flag/block", "Thresholds need calibration for operational deployment"],
        ["Reporting", "openpyxl + ReportLab", "Excel charts and PDF report generated from test run", "Not an independent security audit"],
    ]:
        coverage.append(row)

    for ws in (detail, coverage):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="294766")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in range(1, ws.max_column + 1):
            max_len = max(len(str(ws.cell(r, col).value or "")) for r in range(1, ws.max_row + 1))
            ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 12), 44)
    for col, width in {"A": 35, "B": 22, "D": 32, "E": 16, "G": 16}.items():
        dashboard.column_dimensions[col].width = width
    for ws in (dashboard, detail, coverage):
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def build_pdf_report(results: pd.DataFrame, model_metrics: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.units import mm

    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=landscape(A4),
        leftMargin=10 * mm, rightMargin=10 * mm,
        topMargin=10 * mm, bottomMargin=10 * mm,
        title="PQ-Auth Sentinel Scenario Test Report",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="PQTitle", parent=styles["Title"], alignment=TA_CENTER,
        textColor=colors.HexColor("#14233B"),
    ))

    total = len(results)
    passed = int((results["Test Status"] == "PASS").sum())
    failed = total - passed
    decisions = results["Decision"].value_counts()
    story = [
        Paragraph("PQ-Auth Sentinel — Scenario Test & Approval Report", styles["PQTitle"]),
        Paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]),
        Spacer(1, 4 * mm),
    ]

    summary = [
        ["Measure", "Observed result"],
        ["Scenarios executed", str(total)],
        ["Checks passed / failed", f"{passed} / {failed}"],
        ["Decision distribution", f"ALLOW {int(decisions.get('ALLOW', 0))} · FLAG {int(decisions.get('FLAG', 0))} · BLOCK {int(decisions.get('BLOCK', 0))}"],
        ["Model test accuracy (synthetic dataset)", f"{model_metrics['test_accuracy']:.2%}" if isinstance(model_metrics.get("test_accuracy"), (int, float)) else "Not available"],
        ["Model macro-F1 (synthetic dataset)", f"{model_metrics['test_macro_f1']:.3f}" if isinstance(model_metrics.get("test_macro_f1"), (int, float)) else "Not available"],
    ]
    table = Table(summary, colWidths=[70 * mm, 110 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#294766")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C4D2")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [table, Spacer(1, 4 * mm), Paragraph("Scenario outcomes", styles["Heading2"])]

    outcomes = [["ID", "Scenario", "Signature", "Replay", "ML prediction", "Risk", "Decision", "Test"]]
    for _, row in results.iterrows():
        outcomes.append([
            str(row["Scenario ID"]), str(row["Scenario"]),
            "VALID" if row["Signature Valid"] else "INVALID",
            "DETECTED" if row["Replay Detected"] else "CLEAR",
            str(row["ML Prediction"]), f'{float(row["Risk Score"]):.3f}',
            str(row["Decision"]), str(row["Test Status"]),
        ])
    result_table = Table(
        outcomes, repeatRows=1,
        colWidths=[15 * mm, 43 * mm, 20 * mm, 20 * mm, 37 * mm, 15 * mm, 20 * mm, 17 * mm],
    )
    result_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#294766")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B8C4D2")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    story += [result_table, Spacer(1, 4 * mm), Paragraph("Approval recommendation", styles["Heading2"])]

    if failed == 0:
        recommendation = (
            "CONDITIONAL APPROVAL FOR ACADEMIC DEMONSTRATION ONLY. All checks in this "
            "controlled suite passed. This does not certify production readiness."
        )
    else:
        recommendation = (
            "NOT APPROVED YET. One or more controlled scenario checks failed. Investigate "
            "the failed rows and rerun the suite before requesting academic/demo approval."
        )
    story.append(Paragraph(recommendation, styles["BodyText"]))
    story.append(Spacer(1, 2 * mm))
    story.append(Paragraph(
        "Limitations: behavioural scenarios use controlled feature inputs and the supplied "
        "synthetic dataset. This is not a network-wide monitor, penetration test, or independent "
        "cryptographic audit. Session-only replay state may reset with the application session.",
        styles["Italic"],
    ))
    doc.build(story)
    return output.getvalue()


def render_test_dashboard(bundle: ModelBundle) -> None:
    st.title("Scenario Testing & Reports")
    st.caption("Run the controlled test suite, inspect measured results, and download Excel/PDF reports.")
    st.warning(
        "Academic prototype only. The behavioural scenarios use controlled telemetry inputs; "
        "passing these tests is not production security certification."
    )

    if st.button("Run all scenario tests", type="primary", use_container_width=True):
        try:
            with st.spinner("Running ML-DSA verification and behavioural model tests..."):
                results = run_scenario_suite(bundle)
            st.session_state["scenario_test_results"] = results
            st.session_state["scenario_test_run_at"] = datetime.now(timezone.utc).isoformat()
        except PQCryptoUnavailable as exc:
            st.error(f"ML-DSA tests could not run: {exc}")
            st.info("Check your liboqs installation and confirm ML-DSA-65 is enabled.")
        except Exception as exc:
            st.error(f"The scenario suite failed: {exc}")
            st.exception(exc)

    results = st.session_state.get("scenario_test_results")
    if results is None:
        st.info("Click **Run all scenario tests** to execute the suite and create reports.")
        st.markdown(
            "Included: normal authentication, replay, tampered message, flooding-like behaviour, "
            "spoofing-like behaviour, corrupted signature, and wrong public key."
        )
        return

    total = len(results)
    passed = int((results["Test Status"] == "PASS").sum())
    failed = total - passed
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Scenarios executed", total)
    c2.metric("Checks passed", passed)
    c3.metric("Checks failed", failed)
    c4.metric("Pass rate", f"{passed / total:.1%}" if total else "N/A")

    st.caption(f"Last run: {st.session_state.get('scenario_test_run_at', 'Unknown')}")
    st.subheader("Scenario outcomes")
    st.dataframe(results, use_container_width=True, hide_index=True)

    st.subheader("Risk score by scenario")
    st.bar_chart(results.set_index("Scenario")[["Risk Score"]])

    st.subheader("Decision distribution")
    decision_counts = results["Decision"].value_counts().reindex(["ALLOW", "FLAG", "BLOCK"], fill_value=0)
    st.bar_chart(decision_counts)

    st.subheader("Test status distribution")
    status_counts = results["Test Status"].value_counts().reindex(["PASS", "FAIL"], fill_value=0)
    st.bar_chart(status_counts)

    excel_data = build_excel_report(results, bundle.metrics)
    pdf_data = build_pdf_report(results, bundle.metrics)

    st.subheader("Download reports")
    left, right = st.columns(2)
    with left:
        st.download_button(
            "Download Excel dashboard (.xlsx)",
            data=excel_data,
            file_name="PQ_Auth_Sentinel_Scenario_Dashboard.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
    with right:
        st.download_button(
            "Download PDF approval report",
            data=pdf_data,
            file_name="PQ_Auth_Sentinel_Approval_Report.pdf",
            mime="application/pdf",
            use_container_width=True,
        )
